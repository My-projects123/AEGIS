"""Build the official 6-slide SIH deck. Run after evaluate.py if possible."""

from __future__ import annotations

import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import nsmap
from pptx.util import Emu, Inches, Pt

NAVY = RGBColor(0x0B, 0x1C, 0x2C)
PANEL = RGBColor(0x12, 0x25, 0x36)
GOLD = RGBColor(0xC5, 0xA4, 0x6E)
CYAN = RGBColor(0x3E, 0xC6, 0xC9)
WHITE = RGBColor(0xF4, 0xF7, 0xFA)
MUTED = RGBColor(0xB7, 0xC6, 0xD4)
RED = RGBColor(0xD4, 0x5D, 0x5D)


def _set_run(run, size=16, bold=False, color=WHITE):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def _bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = NAVY


def _rect(slide, l, t, w, h, color):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    return sh


def _text(slide, l, t, w, h, text, size=16, bold=False, color=WHITE, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _set_run(run, size=size, bold=bold, color=color)
    return box


def _bullets(slide, l, t, w, h, items, size=15):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(6)
        run = p.add_run()
        run.text = "▸  " + item
        _set_run(run, size=size, color=WHITE)
    return box


def _chip(slide, l, t, w, h, label, fill):
    sh = _rect(slide, l, t, w, h, fill)
    tf = sh.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = label
    _set_run(run, size=10, bold=True, color=NAVY)
    return sh


def _footer(slide, n):
    _text(slide, Inches(0.4), Inches(7.05), Inches(10), Inches(0.3),
          f"SIH26052  ·  DRDO  ·  AEGIS-COM  ·  Internal hackathon POC  ·  {n}/6",
          size=10, color=MUTED)


def load_eval():
    p = ROOT / "results" / "eval.json"
    if p.exists():
        return json.loads(p.read_text())
    return None


def add_notes(slide, text):
    notes = slide.notes_slide.notes_text_frame
    notes.text = text


def build():
    ev = load_eval()
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # ---- 1 ----
    s = prs.slides.add_slide(blank)
    _bg(s)
    _rect(s, Inches(0), Inches(0), Inches(0.12), Inches(7.5), GOLD)
    _text(s, Inches(0.5), Inches(0.28), Inches(12), Inches(0.35),
          "SIH26052  ·  DRDO  ·  HARDWARE THEME  ·  SMART VEHICLES", 12, False, GOLD)
    _text(s, Inches(0.5), Inches(0.55), Inches(12), Inches(0.55),
          "AEGIS-COM", 36, True, WHITE)
    _text(s, Inches(0.5), Inches(1.1), Inches(12), Inches(0.4),
          "Event-Aware Adaptive Speech Protection for Defence Communications", 18, False, CYAN)
    _bullets(s, Inches(0.5), Inches(1.7), Inches(7.4), Inches(2.8), [
        "Radio commands fail when engine, rotor, and impulsive noise hit the same sentence.",
        "LMS / Wiener / spectral subtraction lag or break on bangs.",
        "We are not showing a fake tactical headset. We show a comms-path controller.",
        "Sense the event → switch enhancement policy → recover the next word.",
    ], size=16)
    _rect(s, Inches(8.2), Inches(1.7), Inches(4.6), Inches(3.3), PANEL)
    _text(s, Inches(8.4), Inches(1.85), Inches(4.2), Inches(0.35), "ONE PIPELINE", 12, True, GOLD)
    for i, (lab, col) in enumerate([
        ("Acoustic intelligence", CYAN),
        ("FSM controller", GOLD),
        ("Mode-specific SE", WHITE),
        ("Recover after impulse", RED),
    ]):
        _chip(s, Inches(8.5), Inches(2.35 + i * 0.55), Inches(4.0), Inches(0.42), lab, col)
    _text(s, Inches(0.5), Inches(5.0), Inches(12.3), Inches(1.6),
          "Laptop POC = software-in-the-loop digital enhancement.  Physical earcup ANC = future companion, not today’s claim.",
          14, False, MUTED)
    _footer(s, 1)
    add_notes(s, "Open with lost radio commands. Name DRDO. Say laptop POC immediately. Do not say we built a headset.")

    # ---- 2 ----
    s = prs.slides.add_slide(blank)
    _bg(s)
    _rect(s, Inches(0), Inches(0), Inches(0.12), Inches(7.5), GOLD)
    _text(s, Inches(0.5), Inches(0.3), Inches(12), Inches(0.5), "Existing systems already protect hearing. The gap is elsewhere.", 26, True, WHITE)
    # two panels
    _rect(s, Inches(0.5), Inches(1.1), Inches(6.0), Inches(5.3), PANEL)
    _text(s, Inches(0.7), Inches(1.25), Inches(5.6), Inches(0.4), "MATURE TODAY", 14, True, GOLD)
    _bullets(s, Inches(0.7), Inches(1.75), Inches(5.6), Inches(4.4), [
        "3M ComTac VI: impulse + steady hearing protection (ANSI S12.42).",
        "Mission Audio Profiles: Comfort / Conversation / Patrol / Observation — operator selected.",
        "INVISIO RA108 / RA5100: PNR + ANR for continuous vehicle noise; user-toggled hear-through.",
        "Generic AI denoisers: one policy for every scene.",
    ], 15)
    _rect(s, Inches(6.8), Inches(1.1), Inches(6.0), Inches(5.3), PANEL)
    _text(s, Inches(7.0), Inches(1.25), Inches(5.6), Inches(0.4), "THE REAL GAP", 14, True, CYAN)
    _bullets(s, Inches(7.0), Inches(1.75), Inches(5.6), Inches(4.4), [
        "MAP is a menu, not a closed loop on the radio path.",
        "ANR is built for rumble, not policy switching after a bang.",
        "AI denoisers smear the next syllable when the noise model is poisoned.",
        "Missing: automatic, confidence-gated enhancement policy + mandatory recovery.",
    ], 15)
    _footer(s, 2)
    add_notes(s, "Never say existing headsets cannot handle impulse. Contrast MAP vs automatic comms policy.")

    # ---- 3 ----
    s = prs.slides.add_slide(blank)
    _bg(s)
    _rect(s, Inches(0), Inches(0), Inches(0.12), Inches(7.5), GOLD)
    _text(s, Inches(0.5), Inches(0.25), Inches(12), Inches(0.45), "Novelty: event-aware policy switching, not ‘AI + DSP’", 24, True, WHITE)
    stages = [
        ("AUDIO", GOLD),
        ("SENSE", CYAN),
        ("DECIDE", GOLD),
        ("PROCESS", CYAN),
        ("RECOVER", RED),
    ]
    for i, (lab, col) in enumerate(stages):
        _chip(s, Inches(0.45 + i * 2.55), Inches(0.95), Inches(2.35), Inches(0.48), lab, col)
    modes = [
        ("SPEECH\nPRESERVE", RGBColor(0x3E, 0xC6, 0xC9)),
        ("STATIONARY", RGBColor(0x7C, 0xB8, 0x7C)),
        ("DYNAMIC", RGBColor(0xE0, 0xB4, 0x4A)),
        ("IMPULSE\nPROTECT", RED),
        ("RECOVERY", GOLD),
    ]
    for i, (lab, col) in enumerate(modes):
        sh = _rect(s, Inches(0.45 + i * 2.55), Inches(1.65), Inches(2.35), Inches(1.15), col)
        tf = sh.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = lab
        _set_run(run, size=13, bold=True, color=NAVY)
    _bullets(s, Inches(0.5), Inches(3.05), Inches(12.3), Inches(3.5), [
        "Every 16 ms we estimate class, severity, speech P, impulse P, confidence — then the FSM chooses a mode.",
        "Impulse has priority and FREEZES the noise estimate. Recovery is mandatory. Low confidence holds the last stable mode.",
        "AI (RNNoise-class on the edge; Wiener in tonight’s POC) is a module inside STATIONARY/DYNAMIC — not the invention.",
        "This is complementary to ComTac MAP / INVISIO ANR, not a claim that those products lack hearing protection.",
    ], 16)
    _footer(s, 3)
    add_notes(s, "Point to RECOVERY. Invention is the mechanism. Enhancer is swappable.")

    # ---- 4 ----
    s = prs.slides.add_slide(blank)
    _bg(s)
    _rect(s, Inches(0), Inches(0), Inches(0.12), Inches(7.5), GOLD)
    _text(s, Inches(0.5), Inches(0.25), Inches(12), Inches(0.45), "Working POC — real processing on a laptop", 24, True, WHITE)
    labels = [
        ("REAL", "Detector, FSM, Wiener, waveforms, spectrograms, SNR, SI-SNR, laptop RTF", CYAN),
        ("SIMULATED", "Formant speech, engine/rotor-like noise, laboratory impulse — not field tapes", GOLD),
        ("TARGET", "SNR > 15 dB · STOI > 0.85 · PESQ > 2.5  (SIH goals, not claimed)", WHITE),
        ("FUTURE", "Jetson / Pi / DSP, dual-mic ANC, tactical headset, DRDO field audio", MUTED),
    ]
    for i, (k, v, c) in enumerate(labels):
        _rect(s, Inches(0.5), Inches(0.9 + i * 1.15), Inches(12.3), Inches(1.05), PANEL)
        _text(s, Inches(0.7), Inches(0.98 + i * 1.15), Inches(2.2), Inches(0.4), k, 16, True, c)
        _text(s, Inches(3.0), Inches(0.98 + i * 1.15), Inches(9.5), Inches(0.8), v, 16, False, WHITE)
    _footer(s, 4)
    add_notes(s, "START DEMO after this slide if time: clean → engine → dynamic → impulse → recovery.")

    # ---- 5 ----
    s = prs.slides.add_slide(blank)
    _bg(s)
    _rect(s, Inches(0), Inches(0), Inches(0.12), Inches(7.5), GOLD)
    _text(s, Inches(0.5), Inches(0.22), Inches(12), Inches(0.45), "Laptop measurements on the synthetic demo scene", 24, True, WHITE)

    def g(path, default="run evaluate.py"):
        if not ev:
            return default
        cur = ev
        for k in path.split("."):
            cur = cur.get(k) if isinstance(cur, dict) else None
            if cur is None:
                return default
        if isinstance(cur, float):
            return f"{cur:.2f}"
        return str(cur)

    demo = (ev or {}).get("demo_scene", {})
    rows = [
        ("Input SNR (dB)", g("demo_scene.input_snr_db")),
        ("Output SNR (dB)", g("demo_scene.output_snr_db")),
        ("SNR improvement (dB)", g("demo_scene.snr_improvement_db")),
        ("SI-SNR in / out (dB)", f"{g('demo_scene.input_si_snr_db')}  →  {g('demo_scene.output_si_snr_db')}"),
        ("Impulse peak reduction (dB)", g("demo_scene.impulse_peak_reduction_db")),
        ("Detect latency / recovery (ms)", f"{g('demo_scene.impulse_detection_latency_ms')}  /  {g('demo_scene.recovery_time_ms')}"),
        ("Laptop RTF", g("demo_scene.rtf")),
    ]
    stoi = demo.get("stoi_enhanced") or {}
    pesq = demo.get("pesq_enhanced") or {}
    stoi_s = f"{stoi.get('value'):.3f} (MEASURED)" if stoi.get("value") is not None else "UNAVAILABLE — TARGET 0.85"
    pesq_s = f"{pesq.get('value'):.3f} (MEASURED)" if isinstance(pesq, dict) and pesq.get("available") else "UNAVAILABLE — TARGET 2.5"
    rows.append(("STOI", stoi_s))
    rows.append(("PESQ", pesq_s))

    for i, (k, v) in enumerate(rows):
        y = Inches(0.75 + i * 0.58)
        _rect(s, Inches(0.5), y, Inches(12.3), Inches(0.52), PANEL if i % 2 == 0 else RGBColor(0x0F, 0x24, 0x34))
        _text(s, Inches(0.7), y + Inches(0.05), Inches(5.5), Inches(0.4), k, 14, False, MUTED)
        _text(s, Inches(6.5), y + Inches(0.05), Inches(6.0), Inches(0.4), str(v), 16, True, CYAN)
    _footer(s, 5)
    add_notes(s, "Read only these numbers. SIH printed thresholds are TARGETS. Hardware was not used.")

    # ---- 6 ----
    s = prs.slides.add_slide(blank)
    _bg(s)
    _rect(s, Inches(0), Inches(0), Inches(0.12), Inches(7.5), GOLD)
    _text(s, Inches(0.5), Inches(0.25), Inches(12), Inches(0.45), "Path to hardware, impact, frozen novelty", 24, True, WHITE)
    _bullets(s, Inches(0.5), Inches(0.9), Inches(12.3), Inches(3.4), [
        "Deploy path: laptop SIL → USB audio → Raspberry Pi / Jetson → boom-mic + certified earcup ANR.",
        "Latency budget (ESTIMATE): 16 ms hop + model + I/O < 60 ms. Laptop RTF is the measured number today.",
        "Impact: fewer lost commands when the scene jumps; software updates without replacing every headset.",
        "Novelty recap: automatic event-aware speech-protection policy with impulse freeze and recovery.",
    ], 16)
    steps = ["POC today", "Edge port", "Headset I/O", "Field trial"]
    for i, lab in enumerate(steps):
        _chip(s, Inches(0.5 + i * 3.15), Inches(4.6), Inches(2.9), Inches(0.55), lab, GOLD if i == 0 else CYAN)
    _text(s, Inches(0.5), Inches(5.4), Inches(12.3), Inches(1.2),
          "We are not asking DRDO to believe a dashboard. We are asking them to believe a controller that changes its mind for a reason — and we showed that reason on audio.",
          16, False, WHITE)
    _footer(s, 6)
    add_notes(s, "Close on Indian software on the comms stack, not replacing ComTac. Invite the demo.")

    out = ROOT / "docs" / "AEGIS-COM_SIH26052.pptx"
    prs.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    build()
